import fitz  # PyMuPDF
from pptx import Presentation
import zipfile
import re
import os
from pathlib import Path


# Security limits to prevent Denial of Service (ZIP bombs, memory exhaustion, decompression bombs)
MAX_FILE_SIZE_BYTES = 20 * 1024 * 1024  # 20MB
MAX_UNCOMPRESSED_ZIP_BYTES = 50 * 1024 * 1024  # 50MB max uncompressed size for PPTX
MAX_ZIP_ENTRIES = 2000
MAX_COMPRESSION_RATIO = 50
MAX_PDF_PAGES = 100
MAX_PPTX_SLIDES = 100
MAX_EXTRACTED_CHARS = 100_000

# Regex patterns for prompt injection detection and neutralization
PROMPT_INJECTION_PATTERNS = [
    re.compile(r"ignore\s+(all\s+)?(previous|prior|above)\s+(instructions|directives|prompts|rules)", re.IGNORECASE),
    re.compile(r"disregard\s+(all\s+)?(previous|prior|above)\s+(instructions|directives|prompts|rules)", re.IGNORECASE),
    re.compile(r"forget\s+(all\s+)?(previous|prior|above)\s+(instructions|directives|prompts|rules)", re.IGNORECASE),
    re.compile(r"you\s+are\s+now\s+(a|an|the)?", re.IGNORECASE),
    re.compile(r"act\s+as\s+(a|an|the)?", re.IGNORECASE),
    re.compile(r"system\s+prompt\s*:", re.IGNORECASE),
    re.compile(r"system\s*override", re.IGNORECASE),
    re.compile(r"developer\s+mode", re.IGNORECASE),
    re.compile(r"jailbreak", re.IGNORECASE),
    re.compile(r"override\s+(all\s+)?instructions", re.IGNORECASE),
    re.compile(r"new\s+instructions\s*:", re.IGNORECASE),
    re.compile(r"respond\s+only\s+with\s+(json|valid\s+json)", re.IGNORECASE),
    re.compile(r"output\s+only\s+score\s*:\s*10", re.IGNORECASE),
]


def sanitize_extracted_text(text: str) -> str:
    """
    Sanitize uploaded file content before injecting into AI prompts.
    Strips prompt injection attempts, system prompt overrides, delimiter hijacking,
    and escapes curly braces so string formatting doesn't fail.
    """
    if not text:
        return ""

    # Remove null bytes and non-printable control characters (except newline, tab, carriage return)
    sanitized = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F\x7F]", "", text)

    # Neutralize prompt injection attempts
    for pattern in PROMPT_INJECTION_PATTERNS:
        sanitized = pattern.sub("[SUSPICIOUS_INSTRUCTION_REMOVED]", sanitized)

    # Neutralize section delimiter hijacking (prevent attacking prompt structure)
    sanitized = sanitized.replace("===", "---")
    sanitized = sanitized.replace("<project_document_content", "<project_doc")
    sanitized = sanitized.replace("</project_document_content>", "</project_doc>")

    # Escape curly braces to avoid breaking Python prompt string templates (.format)
    sanitized = sanitized.replace("{", "(").replace("}", ")")

    # Limit maximum characters from any single file
    if len(sanitized) > MAX_EXTRACTED_CHARS:
        sanitized = sanitized[:MAX_EXTRACTED_CHARS] + "\n\n[Content truncated to prevent resource exhaustion]"

    return sanitized


def validate_file_magic_bytes(file_path: str, expected_type: str) -> None:
    """
    Validate file types server-side using magic bytes inspection,
    ensuring renamed executables, scripts, or malicious files are rejected.
    """
    with open(file_path, "rb") as f:
        header = f.read(4096)

    if not header:
        raise ValueError("Uploaded file is empty.")

    # Check for dangerous executable signatures regardless of reported type
    if header.startswith(b"MZ"):  # Windows PE / EXE / DLL
        raise ValueError("Invalid file content: Executable binaries are not permitted.")
    if header.startswith(b"\x7fELF"):  # Linux ELF executable
        raise ValueError("Invalid file content: Executable binaries are not permitted.")
    if header.startswith(b"\xca\xfe\xba\xbe") or header.startswith(b"\xfe\xed\xfa"):  # Mach-O / Java bytecode
        raise ValueError("Invalid file content: Binary files are not permitted.")

    if expected_type == "pdf":
        # PDF files must start with '%PDF-'
        if not header.startswith(b"%PDF-"):
            raise ValueError("Invalid file content: File does not match valid PDF format.")

    elif expected_type == "pptx":
        # PPTX is an OpenXML ZIP archive starting with PK\x03\x04
        if not (header.startswith(b"PK\x03\x04") or header.startswith(b"PK\x05\x06") or header.startswith(b"PK\x07\x08")):
            raise ValueError("Invalid file content: File does not match valid presentation format.")

        # Inspect ZIP structure and check for ZIP bombs
        try:
            with zipfile.ZipFile(file_path, "r") as zf:
                infolist = zf.infolist()
                if len(infolist) > MAX_ZIP_ENTRIES:
                    raise ValueError("Presentation archive contains too many files (potential decompression bomb).")

                total_uncompressed = 0
                total_compressed = 0
                has_presentation_xml = False
                has_content_types = False

                for info in infolist:
                    # Check for path traversal inside archive
                    if ".." in info.filename or info.filename.startswith("/") or info.filename.startswith("\\"):
                        raise ValueError("Presentation archive contains invalid member paths.")

                    total_uncompressed += info.file_size
                    total_compressed += info.compress_size

                    if info.filename == "[Content_Types].xml":
                        has_content_types = True
                    if info.filename.startswith("ppt/"):
                        has_presentation_xml = True

                # ZIP bomb defense: limit total uncompressed size
                if total_uncompressed > MAX_UNCOMPRESSED_ZIP_BYTES:
                    raise ValueError("Presentation archive exceeds maximum safe uncompressed size.")

                # Compression ratio defense
                if total_compressed > 0 and (total_uncompressed / total_compressed) > MAX_COMPRESSION_RATIO:
                    raise ValueError("Suspicious compression ratio detected (potential ZIP bomb).")

                if not (has_content_types or has_presentation_xml):
                    raise ValueError("Invalid PPTX archive: Missing PowerPoint structure.")

        except zipfile.BadZipFile:
            raise ValueError("Invalid or corrupted PowerPoint file.")

    elif expected_type == "md":
        # Reject if binary or HTML/executable
        if b"\x00" in header:
            raise ValueError("Invalid file content: Binary data detected in markdown file.")
        # Reject HTML files disguised as markdown
        stripped_header = header.strip().lower()
        if stripped_header.startswith(b"<!doctype html") or stripped_header.startswith(b"<html") or b"<script" in stripped_header:
            raise ValueError("Invalid file content: HTML or script code is not permitted in markdown files.")


def extract_pdf_text(file_path: str) -> str:
    """Extract all text from a PDF file with resource limits and safe cleanup."""
    doc = None
    try:
        doc = fitz.open(file_path)
        page_count = len(doc)
        if page_count > MAX_PDF_PAGES:
            raise ValueError(f"PDF exceeds maximum page limit of {MAX_PDF_PAGES} pages.")

        text_parts = []
        total_chars = 0
        for page_num in range(page_count):
            page = doc[page_num]
            text = page.get_text("text")
            if text.strip():
                clean_page_text = text.strip()
                total_chars += len(clean_page_text)
                if total_chars > MAX_EXTRACTED_CHARS:
                    text_parts.append(f"[Page {page_num + 1}]\n{clean_page_text[:1000]}\n[Remaining text truncated]")
                    break
                text_parts.append(f"[Page {page_num + 1}]\n{clean_page_text}")

        return "\n\n".join(text_parts)
    except ValueError:
        raise
    except Exception as e:
        raise ValueError("Failed to parse PDF document. The file may be damaged or invalid.")
    finally:
        # Guarantee closure of PyMuPDF handle so files can be deleted without Windows file locks
        if doc is not None:
            try:
                doc.close()
            except Exception:
                pass


def extract_pptx_text(file_path: str) -> str:
    """Extract all text from a PPTX file with slide count limits."""
    try:
        prs = Presentation(file_path)
        slides = list(prs.slides)
        if len(slides) > MAX_PPTX_SLIDES:
            raise ValueError(f"Presentation exceeds maximum slide limit of {MAX_PPTX_SLIDES} slides.")

        text_parts = []
        total_chars = 0
        for slide_num, slide in enumerate(slides, 1):
            slide_texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_texts.append(shape.text.strip())
            if slide_texts:
                slide_content = "\n".join(slide_texts)
                total_chars += len(slide_content)
                if total_chars > MAX_EXTRACTED_CHARS:
                    text_parts.append(f"[Slide {slide_num}]\n{slide_content[:1000]}\n[Remaining slides truncated]")
                    break
                text_parts.append(f"[Slide {slide_num}]\n" + slide_content)
        return "\n\n".join(text_parts)
    except ValueError:
        raise
    except Exception:
        raise ValueError("Failed to parse presentation file. The file may be damaged or invalid.")


def extract_markdown_text(file_path: str) -> str:
    """Extract and lightly clean markdown content safely."""
    try:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read(MAX_EXTRACTED_CHARS + 1024)

        # Remove HTML tags to prevent XSS / markup abuse
        content = re.sub(r"<[^>]+>", "", content)
        # Clean up image references but keep alt text
        content = re.sub(r"!\[([^\]]*)\]\([^\)]*\)", r"[Image: \1]", content)
        # Keep link text, remove URLs
        content = re.sub(r"\[([^\]]+)\]\([^\)]*\)", r"\1", content)
        # Remove code fence markers but keep content
        content = re.sub(r"```[a-z]*\n?", "", content)
        content = content.replace("```", "")

        return content.strip()
    except ValueError:
        raise
    except Exception:
        raise ValueError("Failed to parse markdown document.")


def parse_uploaded_file(file_path: str, filename: str) -> dict:
    """
    Parse an uploaded file and return sanitized extracted text with metadata.
    Validates file magic bytes, protects against zip/PDF bombs, and sanitizes prompt injections.
    """
    ext = Path(filename).suffix.lower()
    file_size = os.path.getsize(file_path)

    if file_size == 0:
        raise ValueError("Uploaded file is empty.")
    if file_size > MAX_FILE_SIZE_BYTES:
        raise ValueError("Uploaded file exceeds 20MB size limit.")

    if ext == ".pdf":
        validate_file_magic_bytes(file_path, "pdf")
        raw_text = extract_pdf_text(file_path)
        file_type = "pitch_deck_pdf"
    elif ext in (".pptx", ".ppt"):
        validate_file_magic_bytes(file_path, "pptx")
        raw_text = extract_pptx_text(file_path)
        file_type = "pitch_deck_pptx"
    elif ext == ".md":
        validate_file_magic_bytes(file_path, "md")
        raw_text = extract_markdown_text(file_path)
        file_type = "readme"
    else:
        raise ValueError("Unsupported file format. Please upload PDF, PPTX, or .md files.")

    if not raw_text or not raw_text.strip():
        raise ValueError("Could not extract readable text. The document may be image-only, empty, or corrupted.")

    # Sanitize content against prompt injection before returning
    sanitized_text = sanitize_extracted_text(raw_text)

    return {
        "type": file_type,
        "text": sanitized_text,
        "char_count": len(sanitized_text),
    }


def build_project_context(parsed_files: list[dict]) -> str:
    """Combine extracted texts from multiple files into a unified project context."""
    sections = []

    for pf in parsed_files:
        if pf["type"] in ("pitch_deck_pdf", "pitch_deck_pptx"):
            sections.append(f"--- PITCH DECK / PRESENTATION ---\n{pf['text']}")
        elif pf["type"] == "readme":
            sections.append(f"--- PROJECT README ---\n{pf['text']}")

    if not sections:
        raise ValueError("No content could be extracted from the uploaded files.")

    context = "\n\n".join(sections)

    # Truncate to maximum safe context length for AI
    max_chars = 30000
    if len(context) > max_chars:
        context = context[:max_chars] + "\n\n[... content truncated for evaluation ...]"

    return context
